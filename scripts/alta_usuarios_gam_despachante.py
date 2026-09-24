#!/usr/bin/env python3
"""
Da de alta usuarios en GAM llamando por HTTP al procedimiento GeneXus
"Alta_Usuarios_GAM_Excel" (parm: &Name,&UserEmail,&FirstName,&LastName,&password,&rolID).

Lee la pestaña "DESPACHANTE (Rol Despachante)" del Excel de usuarios (columnas
Empresa, CUIT_Empresa, Email, PhoneNumber, UserName, Address, Name, Surname,
Rol, UltimoAcceso, TipoPagoADF, PASSWORD) e invoca el procedimiento una vez por
fila, con rolID=11 (rol Despachante en GAM).

El Excel se pasa como argumento en tiempo de ejecucion y nunca se commitea al
repo, para no dejar contrasenias de usuarios en el historial de git.

Requiere: pip install openpyxl

Uso:
    python3 scripts/alta_usuarios_gam_despachante.py --xlsx /ruta/al/Excel.xlsx [--dry-run]
"""
import argparse
import sys
import time
import urllib.parse
import urllib.request

import openpyxl

BASE_URL = (
    "https://desa02.rendelit.ar/RendelMoveDFWeb17_VDesarrollo.NetEnvironment"
    "/Alta_Usuarios_GAM_Excel.aspx"
)
SHEET_NAME = "DESPACHANTE (Rol Despachante)"
ROL_ID_DESPACHANTE = 11


def load_rows(xlsx_path, sheet_name):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header)}

    result = []
    for r in rows[1:]:
        if not r[idx["UserName"]]:
            continue
        result.append({
            "Empresa": r[idx["Empresa"]],
            "Name": r[idx["UserName"]],
            "UserEmail": r[idx["Email"]],
            "FirstName": r[idx["Name"]],
            "LastName": r[idx["Surname"]],
            "password": r[idx["PASSWORD"]],
            "rolID": ROL_ID_DESPACHANTE,
        })
    return result


def build_url(row):
    params = [
        row["Name"],
        row["UserEmail"],
        row["FirstName"],
        row["LastName"],
        row["password"],
        row["rolID"],
    ]
    encoded = "$".join(urllib.parse.quote(str(p), safe="") for p in params)
    return f"{BASE_URL}?{encoded}"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, help="Ruta al Excel de usuarios")
    parser.add_argument("--sheet", default=SHEET_NAME)
    parser.add_argument("--dry-run", action="store_true", help="Solo imprime las URLs, no llama al servidor")
    parser.add_argument("--delay", type=float, default=0.5, help="Segundos de espera entre llamadas")
    args = parser.parse_args()

    rows = load_rows(args.xlsx, args.sheet)
    print(f"{len(rows)} usuarios a dar de alta desde '{args.sheet}'")

    ok, failed = 0, []
    for row in rows:
        url = build_url(row)
        if args.dry_run:
            print(f"[DRY-RUN] {row['Empresa']} ({row['Name']}) -> {url}")
            continue
        try:
            with urllib.request.urlopen(url, timeout=30) as resp:
                body = resp.read().decode("utf-8", errors="replace")
                print(f"[OK] {row['Empresa']} ({row['Name']}) -> HTTP {resp.status}")
                if body.strip():
                    print(f"     respuesta: {body.strip()[:300]}")
                ok += 1
        except Exception as exc:
            print(f"[ERROR] {row['Empresa']} ({row['Name']}) -> {exc}", file=sys.stderr)
            failed.append(row["Empresa"])
        time.sleep(args.delay)

    if not args.dry_run:
        print(f"\nResumen: {ok} OK, {len(failed)} fallidos")
        if failed:
            print("Fallidos:", ", ".join(failed))


if __name__ == "__main__":
    main()
